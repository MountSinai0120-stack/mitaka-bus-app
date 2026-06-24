# -*- coding: utf-8 -*-
"""3番のりば (武蔵境・営業所・深大寺・65植物公園・66調布) を記入。
IMG_1803(6-12)/1804(7-18)/1805(10-24)。6時の行 = 8 + マーク数(5) = 13。
(分, 前運用route, uncertain)。22〜24時は時刻なし(凡例行)。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PATH = "三鷹駅南口_平日_時刻表.xlsx"
NORM = Font(name="Yu Gothic", size=10)
YELLOW = PatternFill("solid", fgColor="FFF2A8")
ORANGE = PatternFill("solid", fgColor="FCD9A8")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

P3 = {
 6: [(52,"56",False),(59,"",False)],
 7: [(8,"54",False),(21,"51",False),(25,"59",False),(28,"",True),(45,"63",False),(48,"59",False),(54,"",False)],
 8: [(15,"51",False),(19,"52",False),(23,"",True),(35,"",True),(45,"54",False),(53,"55",False),(56,"",True)],
 9: [(15,"54",False),(25,"",True),(45,"54",False),(53,"53",False),(54,"",True)],
 10:[(15,"",False),(28,"53",False),(45,"54",False),(54,"",True)],
 11:[(9,"",True),(15,"56",False),(28,"",True),(37,"54",False),(45,"51",False),(57,"",True)],
 12:[(3,"",True),(15,"54",False),(42,"",True),(45,"",True),(57,"",True)],
 13:[(15,"63",False),(33,"",True),(45,"63",False),(53,"53",True),(57,"",True)],
 14:[(4,"53",False),(15,"",True),(27,"63",False),(45,"63",False),(55,"",True)],
 15:[(15,"61",True),(18,"51",True),(24,"",True),(45,"61",True)],
 16:[(2,"51",True),(23,"",True),(32,"",True),(45,"52",True),(51,"",True),(58,"63",True)],
 17:[(15,"56",True),(22,"",True),(45,"62",True),(48,"59",True),(56,"",True)],
 18:[(11,"60",True),(42,"56",True),(45,"54",True),(50,"",True)],
 19:[(4,"53",True),(15,"",True),(35,"61",True),(48,"",True)],
 20:[(15,"63",True),(18,"",True),(32,"51",True),(45,"",True),(50,"",True)],
 21:[(17,"62",True),(26,"",True),(31,"51",True),(50,"",True)],
}

wb = openpyxl.load_workbook(PATH)
ws = wb["3番のりば"]
HOUR6 = 13
for h, deps in P3.items():
    r = HOUR6 + (h - 6)
    for i, (mm, route, unc) in enumerate(deps):
        cell = ws.cell(r, 2 + i, f"{mm:02d}" + (f"\n鷹{route}" if route else ""))
        cell.font = NORM; cell.alignment = CENTER
        if unc: cell.fill = ORANGE
        elif route: cell.fill = YELLOW
    ws.row_dimensions[r].height = 30
wb.save(PATH)
print("3番のりば 記入完了")
