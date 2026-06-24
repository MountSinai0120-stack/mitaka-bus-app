# -*- coding: utf-8 -*-
"""
2番のりば の時刻を写真(IMG_1794/1795/1796/1797)から best-effort で記入する。
- 各タプル: (分, 前運用route, uncertain)
  前運用route='' なら前運用なし表示。uncertain=True は要確認(オレンジ)。
- 早い時間帯(6〜11時)は確信度高め。12時以降は数字が重なり気味で要確認多め。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PATH = "三鷹駅南口_平日_時刻表.xlsx"

NORM = Font(name="Yu Gothic", size=10)
YELLOW = PatternFill("solid", fgColor="FFF2A8")   # 前運用あり(確信)
ORANGE = PatternFill("solid", fgColor="FCD9A8")    # 要確認
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# hour -> list of (minute, 前運用route, uncertain)
DATA = {
    6:  [(15,"",False),(28,"52",False),(38,"63",False),(48,"51",False),(56,"51",False)],
    7:  [(4,"",False),(8,"63",False),(13,"55",True),(22,"",True),(32,"51",False),(38,"51",False),(50,"60",True)],
    8:  [(3,"51",False),(10,"60",False),(15,"65",True),(34,"55",False),(45,"54",False),(55,"51",True)],
    9:  [(0,"63",False),(15,"59",False),(30,"59",False),(45,"63",True),(50,"63",True)],
    10: [(0,"",False),(15,"51",False),(30,"57",False),(45,"51",True)],
    11: [(0,"",False),(8,"54",False),(23,"53",False),(30,"",True),(45,"63",True),(55,"61",True)],
    12: [(8,"51",True),(13,"",True),(20,"51",True),(27,"",True),(35,"55",True),(42,"",True),(50,"61",True),(57,"",True)],
    13: [(0,"",True),(12,"65",True),(20,"",True),(30,"63",True),(42,"",True),(50,"52",True),(57,"61",True)],
    14: [(9,"55",True),(12,"",True),(20,"52",True),(27,"62",True),(35,"",True),(42,"62",True),(57,"",True)],
    15: [(3,"",True),(12,"54",True),(20,"",True),(35,"54",True),(44,"",True),(50,"63",True)],
    16: [(0,"",True),(12,"51",True),(24,"",True),(38,"",True),(48,"",True),(55,"62",True)],
    17: [(2,"",True),(14,"61",True),(28,"",True),(40,"54",True),(46,"",True),(57,"57",True)],
    18: [(0,"",True),(11,"54",True),(20,"",True),(36,"",True),(48,"62",True),(55,"63",True)],
    19: [(5,"63",True),(11,"",True),(23,"51",True),(35,"",True),(48,"",True),(55,"63",True)],
    20: [(2,"",True),(9,"51",True),(35,"",True),(48,"54",True)],
    21: [(14,"54",True),(23,"",True),(31,"55",True),(53,"",True)],
    22: [(10,"54",True),(40,"",True)],
    23: [(10,"",True)],
    24: [],
}

wb = openpyxl.load_workbook(PATH)
ws = wb["2番のりば"]

# 2番のりば: marks 4行 → table_top = 11, ヘッダ=row11, 6時=row12
TABLE_HEADER = 11
def row_for_hour(h):
    return TABLE_HEADER + 1 + (h - 6)

for h, deps in DATA.items():
    r = row_for_hour(h)
    for i, (mm, route, unc) in enumerate(deps):
        c = 2 + i
        text = f"{mm:02d}"
        if route:
            text += f"\n鷹{route}"
        cell = ws.cell(r, c, text)
        cell.font = NORM
        cell.alignment = CENTER
        if unc:
            cell.fill = ORANGE
        elif route:
            cell.fill = YELLOW

# 24時の注記
ws.cell(row_for_hour(24), 2, "鷹54\n植").font = NORM
ws.cell(row_for_hour(24), 2).alignment = CENTER
ws.cell(row_for_hour(24), 2).fill = ORANGE

# このシートだけ行高を少し上げて2段表示を見やすく
for h in range(6, 25):
    ws.row_dimensions[row_for_hour(h)].height = 30

wb.save(PATH)
print("updated 2番のりば")
