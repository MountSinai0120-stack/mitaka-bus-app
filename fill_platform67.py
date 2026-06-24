# -*- coding: utf-8 -*-
"""6番・7番のりば を写真から best-effort 記入。
table_top(ヘッダ行) = 7 + マーク数。6時の行 = 8 + マーク数。
6番=マーク6個→ヘッダ13/6時14。7番=マーク7個→ヘッダ14/6時15。
(分, 前運用route, uncertain)。uncertain=True→要確認(オレンジ)。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PATH = "三鷹駅南口_平日_時刻表.xlsx"
NORM = Font(name="Yu Gothic", size=10)
YELLOW = PatternFill("solid", fgColor="FFF2A8")
ORANGE = PatternFill("solid", fgColor="FCD9A8")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---- 6番のりば (榊原・朝三・車返・多磨駅・60循環・吉祥寺) ----
P6 = {
 6:[(40,"53",False),(55,"51",False)],
 7:[(9,"59",False),(18,"51",False),(23,"55",False),(29,"",True),(36,"",True),(40,"56",True),(48,"54",True),(51,"65",True)],
 8:[(0,"56",False),(4,"60",False),(8,"56",False),(20,"55",True),(24,"51",True),(37,"56",True),(46,"51",True)],
 9:[(6,"56",False),(26,"53",False),(46,"56",False)],
 10:[(6,"56",False),(31,"",True),(59,"",True)],
 11:[(14,"63",False),(34,"51",False),(46,"62",False)],
 12:[(14,"",True),(34,"",True),(54,"57",True)],
 13:[(9,"57",False),(29,"",True),(47,"54",True)],
 14:[(6,"63",False),(25,"",True),(45,"",True)],
 15:[(5,"",True),(25,"",True),(45,"",True)],
 16:[(5,"53",True),(25,"",True),(45,"63",True),(57,"",True)],
 17:[(12,"51",True),(22,"",True),(27,"",True),(39,"",True),(42,"56",True),(52,"",True),(57,"62",True)],
 18:[(4,"57",True),(16,"51",True),(41,"53",True),(50,"65",True)],
 19:[(4,"54",True),(16,"",True),(40,"56",True)],
 20:[(4,"56",True),(44,"",True)],
 21:[(6,"63",True),(30,"54",True),(55,"",True)],
 22:[(6,"56",True),(18,"",True),(30,"60",True),(45,"",True)],
}

# ---- 7番のりば (仙川・新川団地・61調布・晃華学園東・67/68三鷹駅) ----
# 最繁のりば。フルボード写真では1便あたりの解像度が低く、全体的に低確信→ほぼ要確認。
P7 = {
 6:[(15,"",True),(22,"51",True),(34,"63",True),(54,"51",True)],
 7:[(2,"",True),(8,"",True),(24,"53",True),(32,"63",True),(37,"56",True),(40,"",True),(44,"57",True),(52,"52",True)],
 8:[(5,"",True),(15,"",True),(34,"51",True),(39,"55",True),(46,"54",True),(52,"52",True)],
 9:[(2,"63",True),(13,"",True),(24,"",True),(40,"",True),(52,"57",True)],
 10:[(5,"56",True),(12,"63",True),(18,"",True),(30,"62",True),(44,"",True),(58,"",True)],
 11:[(5,"",True),(24,"",True),(30,"63",True),(46,"",True),(57,"",True)],
 12:[(5,"56",True),(13,"51",True),(28,"",True),(44,"",True),(58,"",True)],
 13:[(5,"",True),(16,"51",True),(28,"",True),(40,"",True),(58,"",True)],
 14:[(5,"",True),(18,"54",True),(28,"",True),(40,"53",True),(52,"",True)],
 15:[(4,"",True),(16,"",True),(28,"",True),(40,"",True),(52,"55",True)],
 16:[(4,"",True),(10,"65",True),(24,"",True),(38,"",True),(52,"",True)],
 17:[(2,"63",True),(14,"",True),(28,"",True),(42,"",True),(54,"",True)],
 18:[(2,"",True),(16,"56",True),(30,"52",True),(44,"",True),(54,"",True)],
 19:[(8,"57",True),(24,"",True),(40,"",True),(52,"",True)],
 20:[(2,"",True),(15,"65",True),(40,"",True)],
 21:[(2,"53",True),(15,"",True),(40,"",True)],
 22:[(21,"51",True)],
 23:[(17,"",True)],
}

def fill(sheet, data, n_marks):
    ws = wb[sheet]
    hour6 = 8 + n_marks
    for h, deps in data.items():
        r = hour6 + (h - 6)
        for i, (mm, route, unc) in enumerate(deps):
            cell = ws.cell(r, 2 + i, f"{mm:02d}" + (f"\n鷹{route}" if route else ""))
            cell.font = NORM; cell.alignment = CENTER
            if unc: cell.fill = ORANGE
            elif route: cell.fill = YELLOW
        ws.row_dimensions[r].height = 30

wb = openpyxl.load_workbook(PATH)
fill("6番のりば", P6, 6)
fill("7番のりば", P7, 7)
wb.save(PATH)
print("updated 6番/7番")
