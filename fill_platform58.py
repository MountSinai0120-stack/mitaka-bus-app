# -*- coding: utf-8 -*-
"""5番・8番のりば を写真から best-effort 記入。
6時の行 = 8 + マーク数。5番=マーク6個→6時14。8番=マーク4個→6時12。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PATH = "三鷹駅南口_平日_時刻表.xlsx"
NORM = Font(name="Yu Gothic", size=10)
YELLOW = PatternFill("solid", fgColor="FFF2A8")
ORANGE = PatternFill("solid", fgColor="FCD9A8")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---- 5番のりば (キリ大・51調布・51飛行場・大沢・十字路・西野) IMG_1809/1810/1811 ----
P5 = {
 6:[(25,"54",False),(40,"",False),(55,"62",False)],
 7:[(10,"51",False),(20,"57",False),(30,"",True),(40,"62",False),(48,"",True),(54,"56",False)],
 8:[(0,"",False),(6,"",False),(12,"61",False),(18,"56",False),(24,"60",True),(30,"61",True),(37,"",True),(45,"59",True),(52,"",True)],
 9:[(0,"61",True),(10,"",True),(20,"65",True),(40,"54",True),(50,"",True)],
 10:[(2,"",True),(14,"65",True),(26,"",True),(38,"",True),(50,"",True)],
 11:[(2,"55",True),(14,"",True),(26,"",True),(38,"",True),(50,"54",True)],
 12:[(2,"65",True),(14,"",True),(26,"61",True),(38,"",True),(50,"",True)],
 13:[(0,"56",True),(10,"",True),(26,"63",True),(30,"",True),(40,"52",True),(50,"",True)],
 14:[(0,"56",True),(10,"",True),(20,"57",True),(30,"",True),(40,"",True),(50,"",True)],
 15:[(0,"56",True),(10,"",True),(20,"",True),(30,"",True),(40,"",True),(50,"",True)],
 16:[(0,"56",True),(10,"",True),(20,"",True),(26,"",True),(30,"",True),(40,"",True)],
 17:[(0,"54",True),(10,"",True),(20,"",True),(30,"",True),(40,"",True),(50,"",True),(58,"61",True)],
 18:[(6,"53",True),(14,"",True),(22,"",True),(30,"",True),(38,"",True),(46,"",True),(54,"55",True)],
 19:[(2,"52",True),(14,"",True),(26,"52",True),(32,"",True),(41,"",True)],
 20:[(6,"52",True),(16,"52",True),(30,"",True),(40,"",True)],
 21:[(6,"52",True),(26,"",True),(36,"61",True)],
 22:[(9,"54",True),(26,"52",True),(52,"",True)],
 23:[(9,"",True),(26,"52",True),(42,"",True)],
}

# ---- 8番のりば (野ヶ谷・杏林井の頭CP・牟礼団地・59循環) IMG_1808/1788 ----
# 写真の解像度が厳しめ→全体的に低確信(要確認)。
P8 = {
 6:[(15,"53",True),(35,"",True),(53,"",True)],
 7:[(0,"52",True),(15,"63",True),(31,"",True),(35,"56",True),(50,"",True)],
 8:[(1,"59",True),(24,"",True),(31,"56",True),(50,"",True)],
 9:[(31,"54",True),(41,"60",True),(48,"56",True)],
 10:[(0,"56",True),(24,"54",True),(30,"51",True),(36,"",True),(47,"",True),(53,"",True)],
 11:[(0,"57",True),(10,"",True),(17,"",True),(33,"",True),(49,"",True)],
 12:[(0,"51",True),(13,"",True),(27,"",True),(42,"56",True)],
 13:[(0,"",True),(13,"52",True),(30,"",True),(45,"",True)],
 14:[(0,"",True),(8,"",True),(11,"61",True),(22,"",True),(28,"",True),(36,"",True)],
 15:[(0,"51",True),(11,"63",True),(26,"",True),(48,"52",True)],
 16:[(0,"57",True),(23,"65",True),(30,"",True),(45,"52",True)],
 17:[(0,"",True),(14,"54",True),(20,"52",True),(46,"",True)],
 18:[(0,"56",True),(5,"",True),(16,"62",True),(42,"",True),(50,"",True)],
 19:[(0,"56",True),(5,"",True),(22,"62",True),(30,"",True)],
 20:[(5,"56",True),(16,"",True),(34,"",True)],
 21:[(6,"52",True),(24,"54",True),(34,"",True),(51,"",True)],
 22:[(10,"52",True)],
}

def fill(sheet, data, n_marks):
    ws = wb[sheet]
    hour6 = 8 + n_marks
    for h, deps in data.items():
        r = hour6 + (h - 6)
        for i, (mm, route, unc) in enumerate(deps):
            cell = ws.cell(r, 2 + i, f"{mm:02d}" + (f"\n鷹{route}" if route else ""))
            cell.font = NORM; cell.alignment = CENTER
            if unc: cell.fill = ORANGE
            elif route: cell.fill = YELLOW
        ws.row_dimensions[r].height = 30

wb = openpyxl.load_workbook(PATH)
fill("5番のりば", P5, 6)
fill("8番のりば", P8, 4)
wb.save(PATH)
print("updated 5番/8番")
